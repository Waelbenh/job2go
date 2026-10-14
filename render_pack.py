"""render_pack.py - turns a content file (JSON) into a CV and a motivation letter.

Usage:  python render_pack.py data/applications/acme_12/content.json

The JSON holds ALL the text (written by Claude Code from profile.yaml).
This script only does the layout. It makes, in the same folder:
  <prefix>_CV.docx / .pdf   and   <prefix>_Letter.docx / .pdf

Two designs (set "design" in the JSON):
  "harvard"  - Times New Roman, centered name, CAPS headings with a dark rule
  "european" - Calibri, name on the left, dark-blue (#1D4E6B) headings with a light rule
"""

import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BLUE = RGBColor(0x1D, 0x4E, 0x6B)
BLACK = RGBColor(0, 0, 0)
GREY = RGBColor(0x55, 0x55, 0x55)
MARGIN_CM = 1.7
TEXT_WIDTH_CM = 21.0 - 2 * MARGIN_CM  # A4 width minus both margins

# Look of each design. Sizes are in points.
DESIGNS = {
    "harvard": dict(font="Times New Roman", body=10.5, name=18, heading=11,
                    accent=BLACK, rule_color="000000", rule_size=8, caps=True, center=True),
    "european": dict(font="Calibri", body=10, name=20, heading=11,
                     accent=BLUE, rule_color="B8CAD6", rule_size=4, caps=False, center=False),
}


# ---------- small helpers ----------

def set_font(run, d, size=None, bold=False, italic=False, color=None):
    run.font.name = d["font"]
    run._element.rPr.rFonts.set(qn("w:eastAsia"), d["font"])  # make Word respect the font
    run.font.size = Pt(size or d["body"])
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def spacing(par, before=0, after=0, line=1.0):
    pf = par.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), line


def bottom_rule(par, color, size):
    """Draw a line under a paragraph."""
    pPr = par._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    line = OxmlElement("w:bottom")
    for key, val in (("val", "single"), ("sz", str(size)), ("space", "1"), ("color", color)):
        line.set(qn(f"w:{key}"), val)
    border.append(line)
    pPr.append(border)


def new_document():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)  # A4
    sec.left_margin = sec.right_margin = Cm(MARGIN_CM)
    sec.top_margin = sec.bottom_margin = Cm(1.5)
    return doc


# ---------- CV ----------

def heading(doc, text, d):
    par = doc.add_paragraph()
    spacing(par, before=7, after=3)
    par.paragraph_format.keep_with_next = True
    run = par.add_run(text.upper() if d["caps"] else text)
    set_font(run, d, d["heading"], bold=True, color=d["accent"])
    bottom_rule(par, d["rule_color"], d["rule_size"])


def entry_line(doc, d, title, org, dates):
    """'Title, Organisation ........ dates' with the dates pushed to the right edge."""
    par = doc.add_paragraph()
    spacing(par, before=2)
    par.paragraph_format.keep_with_next = True
    par.paragraph_format.tab_stops.add_tab_stop(Cm(TEXT_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT)
    set_font(par.add_run(title), d, bold=True)
    if org:
        set_font(par.add_run(f", {org}"), d, italic=(d["font"] == "Times New Roman"))
    if dates:
        set_font(par.add_run(f"\t{dates}"), d, color=GREY)


def bullet(doc, d, text):
    par = doc.add_paragraph(style="List Bullet")
    spacing(par, after=0.5)
    par.paragraph_format.left_indent = Cm(0.55)
    par.paragraph_format.first_line_indent = Cm(-0.4)
    set_font(par.add_run(text), d)


def build_cv(content, d, path):
    cv = content["cv"]
    doc = new_document()
    align = WD_ALIGN_PARAGRAPH.CENTER if d["center"] else WD_ALIGN_PARAGRAPH.LEFT

    # Name, headline, contact line
    par = doc.add_paragraph()
    par.alignment = align
    spacing(par, after=1)
    set_font(par.add_run(cv["name"]), d, d["name"], bold=True, color=d["accent"])
    if cv.get("headline"):
        par = doc.add_paragraph()
        par.alignment = align
        spacing(par, after=1)
        set_font(par.add_run(cv["headline"]), d, d["body"] + 1, bold=(not d["center"]), color=d["accent"])
    par = doc.add_paragraph()
    par.alignment = align
    spacing(par, after=2)
    set_font(par.add_run(cv["contact"]), d, d["body"] - 0.5, color=GREY)

    # Summary
    heading(doc, cv["summary_title"], d)
    par = doc.add_paragraph()
    spacing(par, after=1)
    set_font(par.add_run(cv["summary"]), d)

    # Other sections
    for section in cv["sections"]:
        heading(doc, section["title"], d)
        for entry in section.get("entries", []):
            entry_line(doc, d, entry["title"], entry.get("org", ""), entry.get("dates", ""))
            for text in entry.get("bullets", []):
                bullet(doc, d, text)
        for skill in section.get("skills", []):  # "Label: item, item, item" (no label: just the items)
            par = doc.add_paragraph()
            spacing(par, after=1)
            if skill["label"]:
                set_font(par.add_run(f"{skill['label']}: "), d, bold=True,
                         color=d["accent"] if not d["center"] else None)
            set_font(par.add_run(", ".join(skill["items"])), d)
    doc.save(path)


# ---------- Letter ----------

def build_letter(content, d, path):
    cv, letter = content["cv"], content["letter"]
    doc = new_document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.3)
    sec.top_margin = Cm(2.0)
    body = d["body"] + 1
    align = WD_ALIGN_PARAGRAPH.CENTER if d["center"] else WD_ALIGN_PARAGRAPH.LEFT

    par = doc.add_paragraph()
    par.alignment = align
    spacing(par, after=1)
    set_font(par.add_run(cv["name"]), d, d["name"] - 4, bold=True, color=d["accent"])
    par = doc.add_paragraph()
    par.alignment = align
    spacing(par, after=10)
    set_font(par.add_run(cv["contact"]), d, d["body"] - 0.5, color=GREY)
    bottom_rule(par, d["rule_color"], d["rule_size"])

    def line(text, bold=False, after=0):
        p = doc.add_paragraph()
        spacing(p, after=after, line=1.1)
        set_font(p.add_run(text), d, body, bold=bold)
        return p

    line(letter["date"], after=8)
    for row in letter["recipient"]:
        line(row)
    subject = line(letter["subject"], bold=True)
    subject.paragraph_format.space_before = Pt(10)
    subject.paragraph_format.space_after = Pt(8)

    for text in letter["paragraphs"]:
        p = doc.add_paragraph()
        spacing(p, after=8, line=1.15)
        set_font(p.add_run(text), d, body)
    line(letter["closing"], after=14)
    line(letter["signature"], bold=True)
    doc.save(path)


# ---------- PDF + checks ----------

PDF_TIMEOUT = 60  # seconds one conversion gets
SOFFICE = "/Applications/LibreOffice.app/Contents/MacOS/soffice"


def to_pdf_libreoffice(docx_path):
    """Convert with LibreOffice running invisibly (no Word, so no permission pop-ups)."""
    import subprocess
    import tempfile
    pdf_path = docx_path.with_suffix(".pdf")
    # A private profile folder stops LibreOffice clashing with a copy you have open.
    with tempfile.TemporaryDirectory() as profile:
        result = subprocess.run(
            [SOFFICE, f"-env:UserInstallation=file://{profile}", "--headless",
             "--convert-to", "pdf", "--outdir", str(docx_path.parent), str(docx_path)],
            capture_output=True, text=True, timeout=PDF_TIMEOUT)
    if not pdf_path.exists():
        raise RuntimeError((result.stdout + result.stderr).strip()[-150:])
    return pdf_path


def to_pdf_word(docx_path):
    """Fallback: convert with Microsoft Word (docx2pdf) in a separate process."""
    import subprocess
    pdf_path = docx_path.with_suffix(".pdf")
    result = subprocess.run(
        [sys.executable, "-c", "import sys; from docx2pdf import convert; convert(sys.argv[1], sys.argv[2])",
         str(docx_path), str(pdf_path)],
        capture_output=True, text=True, timeout=PDF_TIMEOUT)
    if not pdf_path.exists():
        raise RuntimeError((result.stdout + result.stderr).strip()[-150:])
    return pdf_path


def to_pdf(docx_path):
    """Make the PDF. LibreOffice first; Word only if LibreOffice is not installed.
    Returns the PDF path, or None if it failed (tries twice)."""
    import subprocess
    convert = to_pdf_libreoffice if Path(SOFFICE).exists() else to_pdf_word
    for attempt in (1, 2):
        try:
            return convert(docx_path)
        except subprocess.TimeoutExpired:
            reason = f"took longer than {PDF_TIMEOUT}s"
        except Exception as error:
            reason = str(error)
        print(f"  PDF conversion failed for {docx_path.name} (try {attempt} of 2): {reason}", flush=True)
    return None


def count_pages(pdf_path):
    return len(re.findall(rb"/Type\s*/Page(?!s)", pdf_path.read_bytes()))


def main():
    if len(sys.argv) != 2:
        sys.exit("Usage: python render_pack.py path/to/content.json")
    content_file = Path(sys.argv[1])
    content = json.loads(content_file.read_text(encoding="utf-8"))
    d = DESIGNS[content["design"]]
    folder = content_file.parent
    prefix = content["file_prefix"]

    # Banned words check (applies to CV and letter text)
    banned = ["spearheaded", "leveraged", "passionate", "dynamic", "results-driven", "synergy",
              "go-getter", "rockstar", "guru", "thrilled"]
    everything = json.dumps(content, ensure_ascii=False).lower()
    found = [w for w in banned if w in everything]
    if found:
        sys.exit(f"Banned words found: {found}")

    words = sum(len(p.split()) for p in content["letter"]["paragraphs"])
    print(f"Letter body: {words} words (target 150-200)")

    cv_docx, letter_docx = folder / f"{prefix}_CV.docx", folder / f"{prefix}_Letter.docx"
    build_cv(content, d, cv_docx)
    build_letter(content, d, letter_docx)
    print(f"Made {cv_docx.name} and {letter_docx.name}")

    for docx_path in (cv_docx, letter_docx):
        pdf = to_pdf(docx_path)
        if pdf:
            pages = count_pages(pdf)
            print(f"  {pdf.name}: {pages} page(s)")
            if docx_path is cv_docx and pages > 1:
                print("  WARNING: the CV is longer than one page - drop the least relevant bullets.")
        else:
            print(f"  Kept {docx_path.name} (no PDF).")


if __name__ == "__main__":
    main()
